#!/usr/bin/env python
# -*- coding: utf-8 -*-
"""
论文数值独立验证脚本  ——  用于在本地终端自行核查，不依赖作者的任何中间脚本
================================================================================

设计原则
--------
本脚本**不复用**论文流水线里的任何派生 CSV。它直接从**特征缓存 .npz** 重算
ProtoNet 的 5-shot 准确率，再与稿件 DOCX 中印出的数字逐格比对。这样即使中间
某一步 CSV 陈旧或写错，也能被独立发现——本次项目正是栽在"派生文件比输入旧"
这一类问题上。

它只依赖：numpy、python-docx、以及项目自身的两个最底层模块
        improve_methods（episode 构造）、calibration（白化与 L2）

用法（在 VSCode 终端中）
------------------------
    & "D:\\soft\\<PY_DIST>\\envs\\<PY_ENV>\\python.exe" _tools\\verify_paper_numbers.py

输出分四段：
    [1] 依赖链时间戳体检 —— 下游文件比上游旧即告警（本次 bug 的直接检测器）
    [2] 输入文件指纹     —— mtime + SHA256，供留痕
    [3] 独立重算         —— 不经 CSV，直接从 .npz 算，与稿件比对
    [4] 结论             —— 逐项 PASS / FAIL / WARN

注意：第 [3] 段只覆盖**声称可复现的部分**。稿件中依赖完整 50 轮 episode 的数字
可与本脚本一致；若某格不一致，**不代表稿件错**，需看该格的数据来源列。
"""
from __future__ import annotations

import hashlib
import os
import sys
import time

import numpy as np
from _paths import find_root, FINAL_OUT, FEAT_CACHE, find_paper  # noqa: E402

ROOT = find_root()
TOOLS = os.path.join(ROOT, "_tools")
FO = FINAL_OUT
CACHE = FEAT_CACHE
DOCX = os.environ.get("PAPER_DOCX") or find_paper()

sys.path.insert(0, TOOLS)

W = 78
report: list[tuple[str, str, str]] = []      # (级别, 项目, 说明)
DOCX_HASH_AT_START = ""
DOCX_MTIME_AT_START = ""


def line(c="-"):
    print(c * W)


def head(t):
    print()
    line("=")
    print(t)
    line("=")


def sha256(p, buf=1 << 20):
    h = hashlib.sha256()
    with open(p, "rb") as f:
        while True:
            b = f.read(buf)
            if not b:
                break
            h.update(b)
    return h.hexdigest()[:16]


# ---------------------------------------------------------------- [1] 时间戳体检
def check_freshness():
    head("[1] 依赖链时间戳体检：下游不得早于上游")
    print("说明：本次论文出现严重错误，根因就是 final_main_table.csv 比特征缓存早 10 小时，")
    print("      而下游脚本仍在复制它的行。此段专门检测这类问题。\n")

    # 上游 -> 下游，按数据真实依赖排序
    CHAIN = [
        ("特征缓存 endb__clean_multitask.npz", os.path.join(CACHE, "endb__clean_multitask.npz")),
        ("基表 final_main_table.csv", os.path.join(FO, "final_main_table.csv")),
        ("校准后 final_main_table_calibrated.csv", os.path.join(FO, "final_main_table_calibrated.csv")),
        ("判据表 discr_projection.csv", os.path.join(FO, "discr_projection.csv")),
        ("固定系数白化 unconditional_whitening.csv", os.path.join(FO, "unconditional_whitening.csv")),
        ("5 路 final_5way_table.csv", os.path.join(FO, "final_5way_table.csv")),
        ("稿件 DOCX", DOCX),
    ]
    prev_name, prev_t = None, None
    for name, p in CHAIN:
        if not os.path.exists(p):
            print("  %-44s 缺失" % name)
            report.append(("FAIL", name, "文件不存在"))
            continue
        t = os.path.getmtime(p)
        mark = ""
        if prev_t is not None and t < prev_t - 60:
            mark = "  <<< 比上游旧 %.1f 小时" % ((prev_t - t) / 3600)
            report.append(("WARN", name, "早于上游 %.1f 小时" % ((prev_t - t) / 3600)))
        print("  %-44s %s%s" % (name, time.strftime("%m-%d %H:%M", time.localtime(t)), mark))
        prev_name, prev_t = name, t
    print("\n  提示：任一 <<< 都意味着下游数据可能陈旧，须重跑该步之后再核对数字。")


# ---------------------------------------------------------------- [2] 输入指纹
def show_fingerprints():
    head("[2] 输入文件指纹（供留痕：mtime + SHA256 前 16 位）")
    files = [DOCX,
             os.path.join(FO, "final_main_table_calibrated.csv"),
             os.path.join(FO, "discr_projection.csv"),
             os.path.join(FO, "final_5way_table.csv"),
             os.path.join(CACHE, "endb__clean_multitask.npz"),
             os.path.join(CACHE, "cub_novel__clean_multitask.npz")]
    for p in files:
        if os.path.exists(p):
            print("  %-46s %s  %s  %8.1f KB"
                  % (os.path.basename(p), time.strftime("%m-%d %H:%M", time.localtime(os.path.getmtime(p))),
                     sha256(p), os.path.getsize(p) / 1024))
        else:
            print("  %-46s 缺失" % os.path.basename(p))


# ------------------------------------------------------------- [3] 独立重算比对
def independent_recompute():
    head("[3] 独立重算：直接从 .npz 特征算 ProtoNet，与稿件逐格比对")

    try:
        from improve_methods import build_episodes, load, EVAL_SEEDS
        import calibration as CAL
    except Exception as e:
        print("  无法导入项目模块：%s" % e)
        report.append(("FAIL", "独立重算", "模块导入失败"))
        return

    try:
        from docx import Document
        doc = Document(DOCX)
    except Exception as e:
        print("  无法读取稿件：%s" % e)
        report.append(("FAIL", "独立重算", "稿件读取失败"))
        return

    print("  协议参数：EVAL_SEEDS = %s ... （共 %d 轮，每轮 5-way 5-shot）"
          % (EVAL_SEEDS[:3], len(EVAL_SEEDS)))
    print("  计算方式：L2 归一化后取类均值原型，最近邻；与本项目 calibration.l2 一致\n")

    def protonet(group, bb, shot=5):
        """直接从特征缓存重算 ProtoNet。不读任何结果 CSV。"""
        f, _, y, _ = load(group, bb)
        idx = build_episodes(y, shot, EVAL_SEEDS)
        F = CAL.l2(f)
        accs = []
        for sup, qry in idx:
            C = np.unique(y[sup])
            P = CAL.l2(np.stack([F[sup][y[sup] == c].mean(0) for c in C]))
            accs.append(100.0 * float((C[np.argmax(CAL.l2(F[qry]) @ P.T, axis=1)] == y[qry]).mean()))
        return float(np.mean(accs)), float(np.std(accs)), len(F), len(idx)

    # 从稿件中取数：表 5 = doc.tables[4]，前两列是数据集与方法名
    t5 = doc.tables[4]
    BBCOL = {2: ("clean_multitask", "领域微调"), 3: ("clean_species", "无属监督"),
             4: ("clip_vitb16", "CLIP B/16"), 5: ("dinov2_vitb14", "DINOv2 B/14")}
    DS = {"cub": "cub_novel", "nab": "nab_novel", "endb": "endb"}
    DSHOW = {"cub": "CUB-200", "nab": "NABirds", "endb": "濒危"}

    print("  核对范围：表 5 中「无校准 / ProtoNet」一行共 12 格（3 数据集 × 4 主干）")
    print("  （这是唯一可用纯特征重算的格子；其余方法含训练或校准步骤，不在此脚本范围）\n")

    cur = None
    checks = []
    for row in t5.rows[1:]:
        c = [x.text.strip() for x in row.cells]
        if c[0] in DS:                      # 表 5 首列为 cub / nab / endb
            cur = c[0]
        if not cur or "ProtoNet" not in c[1] or "无校准" not in c[1]:
            continue
        for ci, (bb, blast) in BBCOL.items():
            checks.append((cur, bb, blast, c[ci]))

    npass = nfail = 0
    print("  %-8s %-14s %10s %14s %10s %s" % ("数据集", "主干", "稿件", "独立重算", "差", "判定"))
    line()
    for ds, bb, blast, cell in checks:
        try:
            m, s, n, nep = protonet(DS[ds], bb)
        except Exception as e:
            print("  %-8s %-14s %10s %14s %10s  重算失败 %s"
                  % (DSHOW[ds], blast, cell, "-", "-", str(e)[:20]))
            report.append(("FAIL", "%s/%s" % (ds, bb), "重算抛异常"))
            nfail += 1
            continue
        try:
            got = float(cell)
        except ValueError:
            print("  %-8s %-14s %10s %14s %10s  稿件该格非数字"
                  % (DSHOW[ds], blast, cell, "%.2f" % m, "-"))
            nfail += 1
            continue
        d = got - m
        ok = abs(d) <= 0.05
        print("  %-8s %-14s %10s %14s %+10.2f  %s   (n=%d 图, %d 轮, ±%.2f)"
              % (DSHOW[ds], blast, cell, "%.2f" % m, d, "PASS" if ok else "FAIL", n, nep, s))
        if ok:
            npass += 1
        else:
            nfail += 1
            report.append(("FAIL", "%s / %s" % (DSHOW[ds], blast),
                           "稿件 %s vs 重算 %.2f（差 %+.2f）" % (cell, m, d)))

    line()
    print("  本节：PASS %d 格，FAIL %d 格" % (npass, nfail))
    if npass + nfail == 0:
        # 匹配失败却照常给出结论，正是本项目反复出现的静默失效
        print("  !! 未匹配到任何单元格——核对未生效，判定为失败而非通过")
        report.append(("FAIL", "独立重算", "未匹配到任何格子，核对未生效"))
    if nfail == 0 and npass > 0:
        report.append(("PASS", "独立重算", "%d 格全部一致" % npass))


# ------------------------------------------------------------------- [4] 结论
def fingerprint_block(nf, nw, np_):
    """验证结论只对某一指纹成立；把该指纹写死在结论里，避免快照失效后被误引。"""
    # 检测已在 check_docx_unchanged() 中做过并计入 nf，此处只读取结果，避免重复计数
    now_hash = sha256(DOCX) if os.path.exists(DOCX) else ""
    changed = bool(DOCX_HASH_AT_START) and now_hash != DOCX_HASH_AT_START

    line("=")
    print("  留痕块（请连同此块一起保存，它界定了本结论的有效范围）")
    line("-")
    print("    验证时间      %s" % time.strftime("%Y-%m-%d %H:%M:%S"))
    print("    稿件路径      %s" % DOCX)
    print("    稿件 mtime    %s" % DOCX_MTIME_AT_START)
    print("    稿件 SHA256   %s" % (DOCX_HASH_AT_START or "缺失"))
    print("    Python        %s / numpy %s" % (sys.version.split()[0], np.__version__))
    print("    判定          %s（失败 %d / 告警 %d / 通过 %d）"
          % ("不通过" if nf else ("有条件通过" if nw else "本次抽检通过"), nf, nw, np_))
    line("-")
    print("    有效范围      仅当稿件 SHA256 仍为 %s 时成立。" % (DOCX_HASH_AT_START or "上值"))
    print("                  稿件一旦重建，指纹即变，本结论自动失效，需重新运行本脚本。")
    if changed:
        print("    !! 本次运行期间稿件已被修改，上述指纹与判定不对应，结论作废。")
    line("=")


def check_docx_unchanged():
    """稿件在本次运行期间是否被修改。必须在打印判定之前调用，
    否则会出现"总判定通过"与留痕块"不通过"自相矛盾。"""
    now = sha256(DOCX) if os.path.exists(DOCX) else ""
    if DOCX_HASH_AT_START and now != DOCX_HASH_AT_START:
        report.append(("FAIL", "稿件在验证期间被修改",
                       "%s -> %s" % (DOCX_HASH_AT_START, now)))
        return True
    return False


def summary():
    # 稿件指纹校验必须在打印判定**之前**执行，否则会出现
    # "[4] 总判定：通过" 与留痕块"不通过"自相矛盾
    check_docx_unchanged()
    head("[4] 结论")
    if not report:
        print("  未产生任何检查项。")
        fingerprint_block(1, 0, 0)
        return
    order = {"FAIL": 0, "WARN": 1, "PASS": 2}
    for lvl, item, note in sorted(report, key=lambda x: order.get(x[0], 3)):
        print("  [%-4s] %-34s %s" % (lvl, item, note))
    nf = sum(1 for r in report if r[0] == "FAIL")
    nw = sum(1 for r in report if r[0] == "WARN")
    np_ = sum(1 for r in report if r[0] == "PASS")
    print()
    line("=")
    if nf:
        print("  总判定：不通过 —— %d 项失败、%d 项告警、%d 项通过" % (nf, nw, np_))
        print("  含义：稿件中存在与本脚本独立重算不符的数字，或依赖链陈旧。")
        print("        请勿在修正前采信稿件数值。")
    elif nw:
        print("  总判定：有条件通过 —— %d 项告警、%d 项通过" % (nw, np_))
        print("  含义：抽检数字一致，但依赖链中存在陈旧文件，未覆盖部分仍可能有问题。")
    else:
        print("  总判定：本次抽检通过 —— %d 项通过" % np_)
        print("  含义：仅表示本脚本覆盖的格子一致。**未覆盖的格子不因此获得保证。**")
    line("=")
    fingerprint_block(nf, nw, np_)
    print()
    print("  本脚本未覆盖（需另行核查）：")
    print("    - 含校准方法族、线性探针的格子（需训练/拟合步骤）")
    print("    - 表 1、表 2、表 4、表 6 与全部图")
    print("    - 参考文献真实性（须用 Crossref / Semantic Scholar 独立核）")
    print("    - 濒危集图像许可与来源（见 image_sources.csv 的 license_verified 列）")
    print("    - 数据集统计量（图像数、属数、观察者数）")
    print()
    print("  提示：本结论绑定于上方留痕块中的稿件指纹。稿件重建后须重新运行。")


def main():
    global DOCX_HASH_AT_START, DOCX_MTIME_AT_START
    DOCX_HASH_AT_START = sha256(DOCX) if os.path.exists(DOCX) else ""
    DOCX_MTIME_AT_START = (time.strftime("%Y-%m-%d %H:%M:%S",
                                         time.localtime(os.path.getmtime(DOCX)))
                           if os.path.exists(DOCX) else "")
    print()
    line("=")
    print("论文数值独立验证  |  %s" % time.strftime("%Y-%m-%d %H:%M:%S"))
    print("稿件：%s" % DOCX)
    print("Python：%s" % sys.version.split()[0])
    print("numpy：%s" % np.__version__)
    line("=")
    check_freshness()
    show_fingerprints()
    independent_recompute()
    summary()


if __name__ == "__main__":
    main()
