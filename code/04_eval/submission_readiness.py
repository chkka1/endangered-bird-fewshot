# -*- coding: utf-8 -*-
"""投稿《计算机工程与应用》前的就绪度核查。

对照 v6 清单「第 3~5 步」与更早的 P1-4（投稿信与声明）。
"""
import io
import os
import re

import hashlib
from docx import Document
from docx.table import Table
from docx.text.paragraph import Paragraph
from docx.shared import Emu
from docx.oxml.ns import qn
from _paths import find_root, find_paper  # noqa: E402

ROOT = find_root()
DOCX = os.environ.get("PAPER_DOCX") or find_paper()
REL = os.path.join(ROOT, "release")

ok_all = True


def head(s):
    print("\n" + "=" * 68)
    print(s)
    print("=" * 68)


d = Document(DOCX)
ps = [x.text for x in d.paragraphs]
t = "\n".join(ps)
sec = d.sections[0]
W = Emu(sec.page_width - sec.left_margin - sec.right_margin).cm
H = Emu(sec.page_height - sec.top_margin - sec.bottom_margin).cm

head("一、v6 P0 四条")
CHK = [
    ("P0-1 英文摘要术语", "unconditional whitening", 0),
    ("P0-2 正文能量值 82.2%", "82.2%", 0),
    ("P0-2 正文能量值 82.4%", "82.4%", 1),
    ("P0-3 六个设定表述", "六个 EfficientNetV2-S 设定", 1),
    ("P0-4 表 3 统计对象说明", "各自 Base 划分的特征上估计", 1),
]
for name, needle, want in CHK:
    got = t.count(needle)
    ok = (got == want) if want in (0, 1) else (got >= want)
    ok_all &= ok
    print("  %-26s %s（%d 处）" % (name, "✓" if ok else "✗", got))

head("二、遗留错误串（应全为 0）")
for k in ("固定的固定", "少数据误判", "少数误判", "审稿标准", "图 3(c)",
          "见表 7", "48.407%", "低四个数量级", "10.38", "12 设定最差"):
    c = t.count(k)
    ok_all &= (c == 0)
    print("  %-16s %d 处 %s" % (k, c, "✓" if c == 0 else "✗"))

head("三、结构与版式")
seq = []
for ch in d.element.body.iterchildren():
    if ch.tag.endswith("}p"):
        seq.append(("P", Paragraph(ch, d).text))
    elif ch.tag.endswith("}tbl"):
        seq.append(("T", Table(ch, d)))
k = 0
cap_ok = True
for i, (kd, v) in enumerate(seq):
    if kd == "T":
        k += 1
        c = seq[i - 1][1].strip() if seq[i - 1][0] == "P" else "***"
        if not c.startswith("表 %d " % k):
            cap_ok = False
wide, tall = [], []
for i, tb in enumerate(d.tables):
    e = tb._tbl.find(qn("w:tblPr")).find(qn("w:tblW"))
    if e is not None and int(e.get(qn("w:w"))) / 567.0 > W + 0.2:
        wide.append(i + 1)
    two = any("\n" in c.text for r in tb.rows for c in r.cells)
    if len(tb.rows) * (0.62 * (2 if two else 1)) > H:
        tall.append(i + 1)
print("  表 %d | 图 %d | 段 %d | 汉字 %d" % (
    k, len(d.inline_shapes), len(ps), len(re.findall(r"[\u4e00-\u9fff]", t))))
print("  版心 %.1f × %.1f cm" % (W, H))
print("  表题逐一对应 : %s" % ("✓" if cap_ok else "✗"))
print("  超宽 %s | 超高 %s" % (wide or "无 ✓", tall or "无 ✓"))
ok_all &= cap_ok and not wide and not tall

head("四、P1 与仓库")
print("  P1-1 108 个 ±s 核验脚本 : %s"
      % ("✓ 有（audit_std.py）" if os.path.exists(
          os.path.join(ROOT, "_tools", "audit_std.py")) else "✗ 无"))
print("  P1-2 仓库地址           : %s"
      % ("✓ 已填" if "github.com" in t or "gitee.com" in t
         else "⏳ 未填（正文写“可向通信作者索取”）"))
print("  P1-3 表 3 单一数据源     : %s"
      % ("✓ 已统一" if "spectrum_metrics.csv" in io.open(
          os.path.join(ROOT, "_tools", "build_paper_v2.py"),
          encoding="utf-8").read() else "✗"))
for f in ("code/README.md", "code/requirements.txt", "results", "paper/v192.docx",
          "LICENSE", "CITATION.cff", ".gitignore"):
    e = os.path.exists(os.path.join(REL, f.replace("/", os.sep)))
    print("  仓库 %-26s %s" % (f, "✓" if e else "✗"))
    ok_all &= e

head("五、仍未完成（P1-4）")
TODO = ["投稿信（cover letter）", "作者贡献声明", "利益冲突声明", "数据可用性声明"]
for x in TODO:
    print("  ⏳ %s" % x)
print("\n  说明：这些是投稿系统通常要求的独立文件，此前几轮清单（v5 的 P1-4）")
print("        曾列出，但一直未写；不写会影响投稿材料的完整性。")

head("结论")
print("  论文与仓库技术项：%s" % ("✓ 就绪" if ok_all else "✗ 仍有缺项"))
print("  投稿材料（信与声明）：⏳ 待补")
